"""A text projection of the Data Extraction Sheet, so the workbook diffs.

The xlsx is the surface humans edit; git can only say "Binary files differ"
about it, which has repeatedly hidden real parameter changes from review.
This module projects every sheet to a CSV in a sibling directory
(``Data Extraction Sheet.d/``) in a canonical form -- data sheets sorted by
their natural key, cell values normalized -- so any workbook change appears
as a row-level text diff, and history questions ("when did this cell
change?") have a `git log -p` answer.

The projection is generated, never hand-edited::

    python -m lsff_utils.extraction_projection            # regenerate
    python -m lsff_utils.extraction_projection --check    # verify in sync

``tests/test_extraction_projection.py`` runs the check, so a commit that
edits the workbook without regenerating fails the fast tests.

Formula cells are projected at their cached value (what pandas sees), and
every formula is additionally listed in ``_formulas.csv`` so an edit to a
formula itself is visible even when its cached value is unchanged.
"""

import argparse
import csv
import filecmp
import io
import re
import sys
import tempfile
from pathlib import Path

import openpyxl

from lsff_utils import paths

WORKBOOK = paths.REPO_ROOT / "0100_data_prep/extraction/Data Extraction Sheet.xlsx"
PROJECTION_DIR = WORKBOOK.with_name(WORKBOOK.stem + ".d")

#: Sort keys for the tabular data sheets (column header names, in precedence
#: order). Sheets not listed (the Progress cross-tabs, Data Sources) are
#: projected in raw row order. Ties beyond the key sort on the remaining
#: columns, so the order is total and deterministic either way.
SHEET_KEYS = {
    "Country-Vehicle Extraction": ["Country", "Vehicle", "Data need", "Data point name", "Quintile", "Sex"],
    "Country-Vehicle-Fort Extraction": ["Country", "Vehicle", "Fortificant", "Data need", "Data point name", "Quintile"],
    "Country Extraction": ["Country", "Data need", "Data point name", "Quintile"],
    "Vehicle Extraction": ["Vehicle", "Data need", "Data point name"],
    "Scenario Definition Extraction": ["Country", "Vehicle", "Fortificant", "Scenario", "Data need", "Data point name"],
    "Universal": ["Data need", "Data point name"],
}

#: Semantic orderings for key columns where alphabetical would scramble
#: related rows apart. Unknown values sort after known ones, alphabetically.
VALUE_ORDER = {
    "Quintile": ["Total", "All (assumed same)", "Lowest", "Second", "Middle", "Fourth", "Highest"],
    "Sex": ["Total", "All (assumed same)", "Female", "Male"],
}


def _render(value):
    """One canonical text form per cell value."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return repr(value)  # shortest round-trip repr, stable across writes
    text = str(value)
    return re.sub(r"\r?\n", r"\\n", text)


def _slug(sheet_name):
    return re.sub(r"[^a-z0-9]+", "_", sheet_name.lower()).strip("_") + ".csv"


def _used_extent(ws):
    """(n_rows, n_cols) of the region that actually holds values."""
    max_row = max_col = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                max_row = max(max_row, cell.row)
                max_col = max(max_col, cell.column)
    return max_row, max_col


def _sort_key(header, keys, row):
    by_name = dict(zip(header, row))
    parts = []
    for col in keys:
        value = by_name.get(col, "")
        order = VALUE_ORDER.get(col)
        if order is not None:
            rank = order.index(value) if value in order else len(order)
            parts.append((rank, value))
        else:
            parts.append((0, value))
    parts.append(tuple(row))  # total order even for duplicate keys
    return parts


def project(workbook_path=WORKBOOK):
    """Return {filename: csv_text} for the whole workbook."""
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    wb_formulas = openpyxl.load_workbook(workbook_path, data_only=False)

    out = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        n_rows, n_cols = _used_extent(ws)
        rows = [
            [_render(ws.cell(row=r, column=c).value) for c in range(1, n_cols + 1)]
            for r in range(1, n_rows + 1)
        ]
        keys = SHEET_KEYS.get(sheet_name)
        if keys and rows:
            header, data = rows[0], rows[1:]
            data.sort(key=lambda row: _sort_key(header, keys, row))
            rows = [header] + data
        buffer = io.StringIO()
        csv.writer(buffer, lineterminator="\n").writerows(rows)
        out[_slug(sheet_name)] = buffer.getvalue()

    formula_rows = [["sheet", "cell", "formula", "cached_value"]]
    for sheet_name in wb_formulas.sheetnames:
        ws = wb_formulas[sheet_name]
        cached = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_rows.append([
                        sheet_name,
                        cell.coordinate,
                        cell.value,
                        _render(cached[cell.coordinate].value),
                    ])
    buffer = io.StringIO()
    csv.writer(buffer, lineterminator="\n").writerows(formula_rows)
    out["_formulas.csv"] = buffer.getvalue()
    return out


def write(projection_dir=PROJECTION_DIR):
    projection_dir.mkdir(exist_ok=True)
    files = project()
    for name, text in files.items():
        (projection_dir / name).write_text(text)
    stale = {p.name for p in projection_dir.glob("*.csv")} - set(files)
    for name in stale:
        (projection_dir / name).unlink()
    return sorted(files)


def check(projection_dir=PROJECTION_DIR):
    """Regenerate to a temp dir and compare; return list of problems."""
    problems = []
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        files = project()
        for name, text in files.items():
            (tmp / name).write_text(text)
            committed = projection_dir / name
            if not committed.exists():
                problems.append(f"missing from projection: {name}")
            elif not filecmp.cmp(committed, tmp / name, shallow=False):
                problems.append(f"out of sync with workbook: {name}")
        for path in projection_dir.glob("*.csv"):
            if path.name not in files:
                problems.append(f"stray file (no matching sheet): {path.name}")
    return problems


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--check", action="store_true",
                        help="verify the committed projection matches the workbook")
    args = parser.parse_args(argv)
    if args.check:
        problems = check()
        for problem in problems:
            print(problem, file=sys.stderr)
        if problems:
            print("regenerate with: python -m lsff_utils.extraction_projection",
                  file=sys.stderr)
            return 1
        print("projection in sync with workbook")
        return 0
    for name in write():
        print(name)
    return 0


if __name__ == "__main__":
    sys.exit(main())
