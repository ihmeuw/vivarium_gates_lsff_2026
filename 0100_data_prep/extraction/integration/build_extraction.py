"""Build the generated extraction tables from the base tables, literature, GF data and decisions.

Inputs (all in this directory unless noted):
    ../data/*.csv    the base extraction tables: hand-entered values with no other source
    lit_long.csv     the literature extraction, from extract_lit.py
    gf_long.csv      the GF background data, from extract_gf.py
    ../../../0050_config/location_fortificant_vehicles.csv
                     the arms the model runs: these take default values and are checked
    decisions.csv    every interpretive choice: values to use, conversions and computation
                     methods, with rationale (see README.md)

Phases, each of which can override the one before:
    1. defaults          for the arms the model runs, from the literature and GF:
                         - measurements (consumption, baseline): literature first, then GF
                           for rows the literature leaves unset;
                         - 2035 targets (fortifiability, intervention coverage and
                           effectiveness): GF first, then the literature.
                         The literature gives a row a value only if its candidates agree (or
                         one is marked 'Use in model'); disagreeing candidates fail the build
                         unless a decision covers the row.
    2. value decisions   active decisions whose source is lit:, value:, gf or keep
                         (times_coverage ones run last, after coverage is final)
    3. method decisions  active decisions whose source is method:<name>; values computed from
                         others already in place (e.g. consumption SDs, CONSUMPTION_DISTRIBUTION.md)
    4. checks            the constraints the pipeline relies on, and unresolved literature
                         candidates (either fails the build)

Methods (how a value is computed) live in this file, in run_method. Which method applies
where, and why, lives in decisions.csv.

Outputs:
    ../generated/*.csv   the tables the data prep notebooks read (never edit by hand)
    changelog.csv    every planned row: base -> new, origin, reference, what it overrode
    review.csv       things a person should look at: GF/literature conflicts, literature
                     values not used, proposed decisions and their effect, placeholders left;
                     every row spelled out (arm, item, values, sources, decision text)
    STATUS.md        the same, per arm, as a readable page, plus the decisions in effect
    comparison.csv   base vs GF vs literature vs output, one row per data row
                     (and comparison.xlsx, the same with colours, when openpyxl is installed)

Usage:
    python build_extraction.py            # write everything
    python build_extraction.py --check    # exit 1 if ../generated/ is out of date
"""

import argparse
import io
import math
import pathlib
import sys

import numpy as np
import pandas as pd

from report import enrich_review, write_status
from common import (
    ALL_SAME, BASE_DIR, MODEL_ARMS_FILE, CV, CVF, DECISIONS_FILE, GENERATED_DIR, GF_LONG, HERE, LIT_LONG,
    NEEDS, QUINTILES, SCEN, TABLES, VEH, Workbook, consumer_moments, format_number, mixture_sd,
    need_alias, norm, round2, table_text,
)

CHANGELOG = HERE / "changelog.csv"
REVIEW = HERE / "review.csv"
STATUS = HERE / "STATUS.md"
COMPARISON = HERE / "comparison.csv"
COMPARISON_XLSX = HERE / "comparison.xlsx"

TARGET_YEAR = "2035"
# Data needs that are 2035 targets rather than measurements. GF is their primary source;
# the literature's values for them describe the current state.
TARGET_NEEDS = {"fortifiability", "intervention_coverage", "intervention_effectiveness"}
# Arms whose consumption, fortifiability and baseline coverage come from HCES microdata
# (0100_data_prep/hces); only their intervention rows take default values.
HCES_ARMS = {("India", "Rice")}
TOTALS_RTOL = 0.1  # check_totals_reasonable in prep_extracted.ipynb
AGREE_RTOL = 0.01  # literature and workbook values closer than this "agree"

SOURCE_LABELS = {
    "lit": "Literature extraction",  # replaced per row with the literature's data source
    "gf": "GF Nutrition PST Background Data (updated 30 Sep 2026)",
    "derived": "DERIVED (MIC-7549)",
    "decision": "decision",  # replaced per decision with a specific label
}


def blank(x):
    return x is None or (isinstance(x, float) and math.isnan(x)) or str(x).strip() == ""


# --------------------------------------------------------------------------------------
# Inputs
# --------------------------------------------------------------------------------------


class GF:
    """Lookups into gf_long.csv."""

    def __init__(self, path):
        self.df = pd.read_csv(path, dtype={"year": str})

    def _rows(self, **criteria):
        d = self.df
        for k, v in criteria.items():
            d = d[d[k].astype(str).str.lower() == str(v).lower()]
        return d

    @staticmethod
    def _out(row):
        value = None if pd.isna(row.value) else float(row.value)
        return value, f"GF '{row.source_sheet}'!{row.source_cell}", row.raw

    def national(self, country, vehicle, metric):
        d = self._rows(source_sheet="Burden x Vehicle Matrix", country=country, vehicle=vehicle, metric=metric)
        if len(d) != 1:
            return None, f"GF: no national {metric} for {country} {vehicle}", None
        return self._out(d.iloc[0])

    def by_quintile(self, country, vehicle, metric):
        d = self._rows(country=country, vehicle=vehicle, metric=metric, stratum_type="wealth_quintile")
        out = {q: (None, f"GF: no {metric} for {q}") for q in QUINTILES}
        for _, row in d.iterrows():
            value, ref, _ = self._out(row)
            out[row.stratum] = (value, ref)
        return out

    def target(self, country, vehicle, metric, year=TARGET_YEAR):
        d = self._rows(source_sheet="Coverage Over Time", country=country, vehicle=vehicle, metric=metric,
                       year=year)
        if len(d) != 1:
            return None, f"GF: no {year} {metric} for {country} {vehicle}", None
        return self._out(d.iloc[0])


def read_arms(wb, model_arms_file=MODEL_ARMS_FILE):
    """Every arm in the base tables, marking the ones the model runs.

    Each arm is a dict: country, vehicle, fortificants (those the model runs, or all in the
    tables for arms it doesn't run), modeled, hces.
    """
    in_tables = {}
    for sheet in (CV, CVF, SCEN):
        for row in wb.rows(sheet):
            key = (wb.get(sheet, row, "country"), wb.get(sheet, row, "vehicle"))
            if None in key:
                continue
            forts = in_tables.setdefault(key, set())
            if sheet != CV and wb.get(sheet, row, "fortificant"):
                forts.add(wb.get(sheet, row, "fortificant"))
    config = pd.read_csv(model_arms_file, dtype=str).fillna("")
    modeled = {}
    for _, a in config.iterrows():
        key = (a.location.strip().title(), a.vehicle.strip().title())
        if key not in in_tables:
            raise ValueError(f"{model_arms_file.name}: {key[0]} {key[1]} has no rows in the extraction tables")
        f = a.fortificant.strip().lower()
        modeled.setdefault(key, set()).update(in_tables[key] if f == "all" else {f.title()})
    return [{"country": c, "vehicle": v, "fortificants": sorted(modeled.get((c, v), forts)),
             "modeled": (c, v) in modeled, "hces": (c, v) in HCES_ARMS}
            for (c, v), forts in sorted(in_tables.items())]


# --------------------------------------------------------------------------------------
# Phase 1: defaults (literature part)
# --------------------------------------------------------------------------------------


def lit_candidates(wb, lit):
    """(table, line) -> the usable literature rows that speak to it."""
    by_row = {}
    for _, r in lit[lit.status == "ok"].iterrows():
        for row in lit_workbook_rows(wb, r):
            by_row.setdefault((NEEDS[r.need][0], row), []).append(r)
    return by_row


def lit_pick(candidates):
    """The literature row to use, or None if the candidates disagree.

    Candidates marked 'Use in model' = no are ignored; if any are marked yes, only those
    count. The rest must all give the same value.
    """
    usable = [x for x in candidates if x.use_in_model != "no"]
    pool = [x for x in usable if x.use_in_model == "yes"] or usable
    if pool and len({round(float(x.value), 9) for x in pool}) == 1:
        return pool[0], [x.lit_id for x in pool]
    return None, []


def default_scope(wb, sheet, row, arms):
    """The arm a row belongs to, if it takes default values (None otherwise)."""
    if sheet == VEH:
        return None  # effect sizes are model-wide; they change only through decisions
    arm = next((a for a in arms if a["modeled"] and a["country"] == wb.get(sheet, row, "country")
                and a["vehicle"] == wb.get(sheet, row, "vehicle")), None)
    if arm is None:
        return None
    if sheet in (CVF, SCEN) and wb.get(sheet, row, "fortificant") not in arm["fortificants"]:
        return None
    if arm["hces"] and sheet != SCEN:
        return None  # consumption, fortifiability and baseline come from HCES for these arms
    return arm


def apply_lit_defaults(wb, lit, arms, gf_values):
    """Phase 1, literature part. Returns {(table, line): candidates} for rows where the
    literature would be the default but its candidates disagree (left unset here)."""
    ambiguous = {}
    for (sheet, row), candidates in sorted(lit_candidates(wb, lit).items(), key=lambda kv: kv[0]):
        if default_scope(wb, sheet, row, arms) is None:
            continue
        if need_alias(wb.get(sheet, row, "data need")) in TARGET_NEEDS and (sheet, row) in gf_values:
            continue  # a 2035 target: GF's value comes first
        x, ids = lit_pick(candidates)
        if x is None:
            ambiguous[(sheet, row)] = candidates
            continue
        pick = " (marked Use in model)" if x.use_in_model == "yes" else ""
        wb.propose(sheet, row, float(x.value), origin="lit",
                   ref=f"{'; '.join(ids)} ({x.source_sheet} row {x.source_row}: {x.data_source})",
                   method=f"literature value as extracted{pick}",
                   source_label=f"Literature extraction: {x.data_source}")
    return ambiguous


# --------------------------------------------------------------------------------------
# Phase 1: defaults (GF part; the mapping documented in GF_DATA_MAPPING.md)
# --------------------------------------------------------------------------------------


def gf_consumption_any(gf, wb, arm):
    c, v = arm["country"], arm["vehicle"]
    national, nat_ref, _ = gf.national(c, v, "coverage")
    by_q = gf.by_quintile(c, v, "coverage")
    method = "GF Coverage (% using the vehicle)"
    if all(val is not None for val, _ in by_q.values()):
        for q, (value, ref) in by_q.items():
            wb.propose(CV, wb.one(CV, country=c, vehicle=v, need="any", quintile=q), value,
                       origin="gf", ref=ref, method=method)
        wb.propose(CV, wb.one(CV, country=c, vehicle=v, need="any", quintile="Total"), national,
                   origin="gf", ref=nat_ref, method=method + "; Total only sanity-checks the quintiles")
        u5_rows = wb.find(CV, country=c, vehicle=v, need="u5_any")
        if v == "Wheat" and u5_rows:  # unused by the pipeline; keep wheat's mirror of WRA tidy
            for row in u5_rows:
                q = wb.get(CV, row, "quintile")
                value, ref = by_q[q] if q in by_q else (national, nat_ref)
                wb.propose(CV, row, value, origin="gf", ref=ref, method=method + "; U5 row mirrors WRA")
    elif national is not None:
        wb.propose(CV, wb.one(CV, country=c, vehicle=v, need="any", quintile=ALL_SAME), national,
                   origin="gf", ref=nat_ref, method=method + "; national, no quintile breakdown in GF")


def gf_consumption_amount(gf, wb, arm):
    """GF g/cap as published, as the mean over all WRA. Conversions are decisions."""
    c, v = arm["country"], arm["vehicle"]
    national, nat_ref, _ = gf.national(c, v, "g_per_capita")
    by_q = gf.by_quintile(c, v, "g_per_capita")
    if any(gcap is None for gcap, _ in by_q.values()):
        return
    for q, (gcap, ref) in by_q.items():
        row = wb.one(CV, country=c, vehicle=v, need="amount", data_point_name="mean", quintile=q)
        wb.propose(CV, row, gcap, origin="gf", ref=ref, method="GF g/cap as published (mean over all WRA)")
    # AUTOMATIC RULE (listed in STATUS.md and GF_DATA_MAPPING.md): if GF's national g/cap is
    # more than TOTALS_RTOL from the mean of its own quintiles, keep the base table's
    # Total instead, because prep_extracted's totals check would otherwise fail. A decision
    # can override this (e.g. method:scale_to_gf_total).
    row = wb.one(CV, country=c, vehicle=v, need="amount", data_point_name="mean", quintile="Total")
    q_mean = np.mean([gcap for gcap, _ in by_q.values()])
    if np.isclose(q_mean, national, rtol=TOTALS_RTOL, atol=0):
        wb.propose(CV, row, national, origin="gf", ref=nat_ref,
                   method="GF national g/cap as published; only sanity-checks the quintiles")
    else:
        wb.keep(CV, row, origin="gf", ref=nat_ref,
                method=f"GF national g/cap {national:.3g} differs from the quintile mean {q_mean:.3g} "
                       "(GF national is from M4N, strata from NFCMS); kept the base Total so the "
                       "pipeline's totals check passes")


def gf_fortifiability(gf, wb, arm):
    c, v = arm["country"], arm["vehicle"]
    value, ref, raw = gf.target(c, v, "consolidation")
    rows = wb.find(CV, country=c, vehicle=v, need="fortifiability")
    if arm["hces"]:
        for row in rows:
            wb.keep(CV, row, origin="gf", ref=ref,
                    method=f"GF {TARGET_YEAR} consolidation {raw} is market-wide; ours applies only to purchased "
                           "non-PDS rice. Needs a decision")
        return
    if value is None:
        return
    for row in rows:
        wb.propose(CV, row, value, origin="gf", ref=ref,
                   method=f"GF {TARGET_YEAR} consolidation; national, applied to every quintile")


def gf_baseline(gf, wb, arm, fortificant):
    c, v = arm["country"], arm["vehicle"]
    any_rows = wb.find(CVF, country=c, vehicle=v, fortificant=fortificant, need="baseline_any")
    eff_rows = wb.find(CVF, country=c, vehicle=v, fortificant=fortificant, need="baseline_effectiveness")
    cons, cons_ref, _ = gf.national(c, v, "consolidation")
    compl, compl_ref, _ = gf.national(c, v, "compliance")
    if arm["hces"]:
        for row in eff_rows:
            wb.keep(CVF, row, origin="gf", ref=compl_ref,
                    method="India baseline coverage comes from HCES; GF's market-wide numbers don't map onto it")
        return
    if cons is None or compl is None or len(any_rows) != 1:
        return  # GF n/a (or quintile-specific rows we don't overwrite with a national number)
    split = round2(math.sqrt(compl))
    coverage = cons * split
    wb.propose(CVF, any_rows[0], coverage, origin="gf", ref=f"{cons_ref} x {compl_ref}",
               method=f"consolidation x sqrt(compliance) = {cons} x {split}")
    if coverage > 0 and len(eff_rows) == 1:
        wb.propose(CVF, eff_rows[0], split, origin="gf", ref=compl_ref,
                   method=f"sqrt(current compliance {compl}), 2 d.p.")


def gf_intervention(gf, wb, arm, fortificant):
    c, v = arm["country"], arm["vehicle"]
    compl, ref, raw = gf.target(c, v, "compliance")
    if compl is None:
        return
    split = round2(math.sqrt(compl))
    for need in ["intervention_coverage", "intervention_effectiveness"]:
        for row in wb.find(SCEN, country=c, vehicle=v, fortificant=fortificant, need=need):
            wb.propose(SCEN, row, split, origin="gf", ref=ref,
                       method=f"sqrt({TARGET_YEAR} compliance {compl}), 2 d.p. (equal split, KB 9/9)")


def apply_gf_defaults(gf, wb, arms):
    for arm in arms:
        if not arm["modeled"]:
            continue
        if not arm["hces"]:
            gf_consumption_any(gf, wb, arm)
            gf_consumption_amount(gf, wb, arm)
        gf_fortifiability(gf, wb, arm)
        for fortificant in arm["fortificants"]:
            gf_baseline(gf, wb, arm, fortificant)
            gf_intervention(gf, wb, arm, fortificant)


# --------------------------------------------------------------------------------------
# Phases 2 and 3: decisions
# --------------------------------------------------------------------------------------

DECISION_KEYS = ["country", "vehicle", "fortificant", "scenario", "quintile", "sex", "data_point_name"]


def decision_rows(wb, d):
    alias = need_alias(d["need"])
    if alias is None:
        raise ValueError(f"decision {d['decision_id']}: unknown need '{d['need']}'")
    sheet = NEEDS[alias][0]
    criteria = {k: d.get(k) for k in DECISION_KEYS if not blank(d.get(k))}
    if sheet == VEH:
        criteria = {k: v for k, v in criteria.items() if k in ("vehicle", "data_point_name")}
    return sheet, wb.find(sheet, need=alias, **criteria)


def coverage_row(wb, sheet, row):
    """The WRA coverage row that goes with an amount row (national for U5 rows)."""
    need = need_alias(wb.get(sheet, row, "data need"))
    q = wb.get(sheet, row, "quintile") if need == "amount" else "Total"
    return wb.one(CV, country=wb.get(sheet, row, "country"), vehicle=wb.get(sheet, row, "vehicle"),
                  need="any", quintile=q)


def likely_lit_rows(lit, d):
    """Literature rows matching a decision's target, to suggest when its lit ID went missing."""
    c = lit[lit.need == need_alias(d["need"])]
    for col in ["country", "vehicle", "fortificant", "quintile", "sex", "data_point_name"]:
        if not blank(d.get(col)):
            c = c[c[col].fillna("").str.lower() == str(d[col]).lower()]
    if c.empty:
        return "none found"
    return "; ".join(f"{r.lit_id} = {r.raw_value} {r.raw_units or ''} from {r.data_source} "
                     f"('{r.source_sheet}' row {r.source_row})" for _, r in c.iterrows())


def resolve_decision(wb, lit, d, gf_values):
    """What a decision would do: a list of effects, one per targeted row.

    Each effect is a dict with sheet, row, action (set / keep / skip), value, derivation,
    origin, ref, method, label (Data source text) and source (for the changelog). 'keep'
    restores the base value; 'skip' leaves whatever earlier phases set.
    """
    sheet, rows = decision_rows(wb, d)
    if not rows:
        raise ValueError(f"decision {d['decision_id']} matches no rows in the extraction tables")
    source, transform = str(d["source"]).strip(), norm(d.get("transform"))
    rationale = d["rationale"]
    out = []
    for row in rows:
        effect = {"sheet": sheet, "row": row, "derivation": None, "origin": "decision", "method": rationale}
        if source == "keep":
            out.append({**effect, "action": "keep", "value": None, "ref": "decision: keep the base value"})
            continue
        if source.startswith("method:"):
            out.append({**effect, **run_method(wb, lit, source.split(":", 1)[1], sheet, row, d)})
            continue
        if source == "gf":
            if (sheet, row) not in gf_values:
                out.append({**effect, "action": "skip", "value": None,
                            "ref": "decision: GF has no value for this row; left as it was"})
                continue
            value, ref = gf_values[(sheet, row)]
            label, kind = SOURCE_LABELS["gf"], "GF"
        elif source.startswith("lit:"):
            match = lit[lit.lit_id == source]
            if len(match) != 1:
                raise ValueError(f"decision {d['decision_id']}: literature row {source} not found. It probably "
                                 "changed in the literature workbook; rows that could replace it: "
                                 + likely_lit_rows(lit, d))
            m = match.iloc[0]
            if pd.isna(m.value):
                raise ValueError(f"decision {d['decision_id']}: {source} has no value")
            value = float(m.value)
            ref = f"{source} ({m.source_sheet} row {m.source_row}: {m.data_source})"
            label, kind = f"Literature extraction: {m.data_source}", "literature"
        elif source.startswith("value:"):
            value = float(source.split(":", 1)[1])
            ref, label, kind = f"value typed in decision {d['decision_id']}", f"Decision {d['decision_id']}", "typed value"
        else:
            raise ValueError(f"decision {d['decision_id']}: unknown source '{source}'")
        derivation = None
        if transform == "sqrt":
            derivation = f"Value = sqrt({format_number(value)}), 2 d.p."
            value = round2(math.sqrt(value))
            rationale_m = f"sqrt, 2 d.p.; {rationale}"
        elif transform == "times_coverage":
            any_row = coverage_row(wb, sheet, row)
            coverage = wb.current(CV, any_row)
            key = ", ".join(str(wb.get(CV, any_row, c)) for c in ["vehicle", "quintile", "sex", "data need"])
            derivation = f"Value = {format_number(value)}*{format_number(coverage)} [Value of: {key}]"
            value = value * coverage
            rationale_m = f"x WRA coverage; {rationale}"
            kind += " x coverage"
        elif transform:
            raise ValueError(f"decision {d['decision_id']}: unknown transform '{transform}'")
        else:
            rationale_m = rationale
        out.append({**effect, "action": "set", "value": value, "derivation": derivation, "ref": ref,
                    "method": rationale_m, "label": label, "source": kind})
    return out


# --------------------------------------------------------------------------------------
# Methods: how a value is computed from others (see CONSUMPTION_DISTRIBUTION.md)
# --------------------------------------------------------------------------------------

WOMEN_POPULATIONS = {"non-pregnant wra", "wra", "all", ""}


def lit_one(lit, *, country, vehicle, need, point, quintile=None, sex=None):
    """The single literature row for a consumption statistic ('Use in model' pick if several)."""
    d = lit[(lit.status == "ok") & (lit.country == country) & (lit.vehicle == vehicle) & (lit.need == need)
            & (lit.data_point_name == point)]
    if need == "amount":
        d = d[d.population.fillna("").str.lower().isin(WOMEN_POPULATIONS) & (d.quintile == quintile)]
    else:
        d = d[d.sex == sex]
    if len(d) > 1 and (d.use_in_model == "yes").sum() == 1:
        d = d[d.use_in_model == "yes"]
    if len(d) != 1:
        raise ValueError(f"expected one literature {point} for {country} {vehicle} {need} "
                         f"{quintile or sex}, found {len(d)} ({', '.join(d.lit_id)})")
    return d.iloc[0]


def matching_row(wb, sheet, row, **changes):
    """The row with the same keys as `row`, except for `changes` (e.g. the mean for an SD row)."""
    keys = {k: wb.get(sheet, row, k.replace("_", " ")) for k in ["country", "vehicle", "quintile", "sex"]}
    keys["need"] = need_alias(wb.get(sheet, row, "data need"))
    keys["data_point_name"] = wb.get(sheet, row, "data point name")
    keys.update(changes)
    return wb.one(sheet, **keys)


def method_consumer_cv(wb, lit, sheet, row, d):
    """SD over all women when consumers get the literature's coefficient of variation.

    The literature (NFCMS usual intake) gives a mean m and SD s over everyone, with no zero
    mass. Our model has a zero mass 1 - p, so it can't match both. Keep the workbook mean
    and give consumers the literature CV s/m:
        mu = m / p,  sigma = CV mu,  SD over all = sqrt(p sigma^2 + p (1 - p) mu^2)
    """
    need = need_alias(wb.get(sheet, row, "data need"))
    if need not in ("amount", "u5_amount") or norm(wb.get(sheet, row, "data point name")) != "standard deviation":
        raise ValueError(f"decision {d['decision_id']}: method consumer_cv applies only to amount SD rows")
    c, v = wb.get(sheet, row, "country"), wb.get(sheet, row, "vehicle")
    key = {"quintile": wb.get(sheet, row, "quintile")} if need == "amount" else {"sex": wb.get(sheet, row, "sex")}
    lit_mean = lit_one(lit, country=c, vehicle=v, need=need, point="mean", **key)
    lit_sd = lit_one(lit, country=c, vehicle=v, need=need, point="standard deviation", **key)
    cv = float(lit_sd.value) / float(lit_mean.value)
    m = wb.current(sheet, matching_row(wb, sheet, row, data_point_name="mean"))
    p = wb.current(CV, coverage_row(wb, sheet, row))
    who = "women" if need == "amount" else "children (national WRA coverage)"
    return {"action": "set", "value": mixture_sd(p, m, cv),
            "ref": f"{lit_sd.lit_id} / {lit_mean.lit_id} ({lit_sd.data_source})",
            "label": f"Derived from {lit_sd.data_source}: literature CV applied to consumers",
            "source": "literature CV applied to consumers",
            "method": f"consumer CV {cv:.3f} (= {float(lit_sd.value):.4g} / {float(lit_mean.value):.4g}); "
                      f"consumer mean = {m:.4g} / coverage {p:.3g}; SD over all {who} = "
                      f"sqrt(p (CV mu)^2 + p (1 - p) mu^2); {d['rationale']}"}


def method_derive_from(wb, lit, sheet, row, d, base):
    """Placeholder from another vehicle: its consumer CV (SD rows) or U5/WRA ratio (U5 means)."""
    need = need_alias(wb.get(sheet, row, "data need"))
    point = norm(wb.get(sheet, row, "data point name"))
    v = wb.get(sheet, row, "vehicle")
    base_row = matching_row(wb, sheet, row, vehicle=base)
    common = {"action": "set", "origin": "derived", "label": SOURCE_LABELS["derived"], "source": f"derived from {base}"}
    if point == "standard deviation":
        mean_a, var_a = consumer_moments(wb.current(CV, coverage_row(wb, sheet, base_row)),
                                         wb.current(sheet, matching_row(wb, sheet, base_row, data_point_name="mean")),
                                         wb.current(sheet, base_row))
        cv = math.sqrt(var_a) / mean_a
        p = wb.current(CV, coverage_row(wb, sheet, row))
        m = wb.current(sheet, matching_row(wb, sheet, row, data_point_name="mean"))
        return {**common, "value": mixture_sd(p, m, cv), "ref": f"(no SD source) {base} consumer CV",
                "method": f"DERIVED: {base} consumer CV {cv:.3f}, converted with coverage {p}; {d['rationale']}"}
    if need == "u5_amount" and point == "mean":
        def wra_quintile_mean(vehicle):
            return np.mean([wb.current(CV, wb.one(CV, country=wb.get(sheet, row, "country"), vehicle=vehicle,
                                                  need="amount", data_point_name="mean", quintile=q))
                            for q in QUINTILES])
        ratio = wra_quintile_mean(v) / wra_quintile_mean(base)
        return {**common, "value": wb.current(sheet, base_row) * ratio, "ref": f"(no U5 source) {base} U5 x {ratio:.4f}",
                "method": f"DERIVED: {base} U5 mean x {ratio:.4f} (ratio of mean WRA quintile means); {d['rationale']}"}
    raise ValueError(f"decision {d['decision_id']}: method derive_from applies to amount SD rows and U5 mean rows")


GF_DATA = None  # the GF lookup, set in main(); methods that need GF's numbers use it


def method_scale_to_gf_total(wb, lit, sheet, row, d):
    """Rescale an arm's amount means so they agree with GF's national g/cap.

    For when GF's national figure is newer than its quintile figures. Every quintile is
    multiplied by f = GF national / mean(GF quintiles), keeping the wealth gradient.
    Quintiles are population fifths, so their unweighted mean is the national mean. The
    Total row is set to the same scaled mean, on whatever basis the rows are already on
    (e.g. x coverage if a times_coverage decision applies), so the pipeline's totals
    check passes.
    """
    need = need_alias(wb.get(sheet, row, "data need"))
    if need != "amount" or norm(wb.get(sheet, row, "data point name")) != "mean":
        raise ValueError(f"decision {d['decision_id']}: method scale_to_gf_total applies only to amount mean rows")
    c, v = wb.get(sheet, row, "country"), wb.get(sheet, row, "vehicle")
    national, nat_ref, _ = GF_DATA.national(c, v, "g_per_capita")
    gf_q = [val for val, _ in GF_DATA.by_quintile(c, v, "g_per_capita").values()]
    if national is None or any(x is None for x in gf_q):
        raise ValueError(f"decision {d['decision_id']}: GF has no national and quintile g/cap for {c} {v}")
    f = national / np.mean(gf_q)
    q_rows = [wb.one(CV, country=c, vehicle=v, need="amount", data_point_name="mean", quintile=q) for q in QUINTILES]
    current_q_mean = np.mean([wb.current(CV, r) for r in q_rows])
    q = wb.get(sheet, row, "quintile")
    value = f * current_q_mean if q == "Total" else f * wb.current(sheet, row)
    return {"action": "set", "value": value, "ref": f"{nat_ref} / mean of GF quintile g/cap",
            "label": f"{SOURCE_LABELS['gf']}: quintiles rescaled to the national g/cap",
            "source": "GF, rescaled to GF national",
            "method": f"x {f:.4f} (GF national {national:g} / GF quintile mean {np.mean(gf_q):.4g}); "
                      f"{d['rationale']}"}


def run_method(wb, lit, name, sheet, row, d):
    if name == "consumer_cv":
        return method_consumer_cv(wb, lit, sheet, row, d)
    if name == "scale_to_gf_total":
        return method_scale_to_gf_total(wb, lit, sheet, row, d)
    if name.startswith("derive_from:"):
        return method_derive_from(wb, lit, sheet, row, d, name.split(":", 1)[1].strip())
    raise ValueError(f"decision {d['decision_id']}: unknown method '{name}' (known: consumer_cv, scale_to_gf_total, derive_from:<vehicle>)")


def is_method(d):
    return str(d["source"]).strip().startswith("method:")


def apply_decisions(wb, lit, decisions, gf_values):
    """Phase 2 (values; times_coverage last) then phase 3 (methods, in file order)."""
    decided = {}  # (sheet, row) -> decision_id, for silencing conflicts
    active = decisions[decisions.status.str.lower() == "active"]
    methods = active[active.apply(is_method, axis=1)]
    values = active[~active.apply(is_method, axis=1)]
    ordered = pd.concat([values[values["transform"].str.lower() != "times_coverage"],
                         values[values["transform"].str.lower() == "times_coverage"], methods])
    for _, d in ordered.iterrows():
        effects = resolve_decision(wb, lit, d, gf_values)  # all rows resolved before any is applied
        for e in effects:
            if e["action"] == "skip":
                continue  # the decision doesn't apply to this row (e.g. GF has no value)
            decided[(e["sheet"], e["row"])] = d.decision_id
            if e["action"] == "set":
                wb.propose(e["sheet"], e["row"], e["value"], origin=e["origin"], ref=e["ref"], method=e["method"],
                           derivation=e.get("derivation"), decision_id=d.decision_id, source_label=e["label"],
                           source=e["source"])
            elif e["action"] == "keep":
                wb.keep(e["sheet"], e["row"], origin="decision", ref=e["ref"], method=e["method"],
                        decision_id=d.decision_id)
    return decided


def among_consumers_arms(decisions):
    """Arms whose amount means are converted with times_coverage (for comparisons)."""
    active = decisions[(decisions.status.str.lower() == "active")
                       & (decisions["transform"].str.lower() == "times_coverage")]
    return {(d.country, d.vehicle) for _, d in active.iterrows() if need_alias(d.need) == "amount"}


# --------------------------------------------------------------------------------------
# Phase 4: checks
# --------------------------------------------------------------------------------------


def run_checks(wb, arms, unresolved=None):
    """Problems that stop the build: broken constraints, and unresolved literature candidates."""
    problems = []
    for (sheet, row), candidates in sorted((unresolved or {}).items()):
        label = " ".join(str(v) for v in wb.describe(sheet, row).values() if not blank(v))
        problems.append(f"{sheet} row {row} ({label}): the literature has disagreeing candidates "
                        + ", ".join(f"{x.lit_id} = {x.value:g} ({x.data_source})" for x in candidates)
                        + ". Mark one 'Use in model' in the literature workbook, or add a decision")
    for (sheet, row), entry in wb.plan.items():
        if entry["action"] == "set" and norm(wb.get(sheet, row, "units")) == "%" and not 0 <= entry["value"] <= 1:
            problems.append(f"{sheet} row {row}: percentage {entry['value']} outside [0, 1]")

    for arm in arms:
        c, v = arm["country"], arm["vehicle"]
        if not arm["modeled"] or arm["hces"] or not wb.find(CV, country=c, vehicle=v, need="amount"):
            continue
        # Pregnancy sim: consumers' variance must be positive (intervention.py)
        for row in wb.find(CV, country=c, vehicle=v, need="amount", data_point_name="mean"):
            q = wb.get(CV, row, "quintile")
            sd_rows = (wb.find(CV, country=c, vehicle=v, need="amount", data_point_name="standard deviation",
                               quintile=q)
                       or wb.find(CV, country=c, vehicle=v, need="amount", data_point_name="standard deviation",
                                  quintile=ALL_SAME))
            any_rows = (wb.find(CV, country=c, vehicle=v, need="any", quintile=q)
                        or wb.find(CV, country=c, vehicle=v, need="any", quintile=ALL_SAME))
            if not sd_rows or not any_rows:
                continue
            _, var_a = consumer_moments(wb.current(CV, any_rows[0]), wb.current(CV, row), wb.current(CV, sd_rows[0]))
            if var_a <= 0:
                problems.append(f"{c} {v} {q}: consumers' variance {var_a:.3g} <= 0 (SD too small for the mean "
                                "and coverage; the pregnancy sim would draw NaN)")
        # prep_extracted: Total within 10% of the unweighted mean of the quintiles
        q_rows = [wb.find(CV, country=c, vehicle=v, need="amount", data_point_name="mean", quintile=q) for q in QUINTILES]
        t_rows = wb.find(CV, country=c, vehicle=v, need="amount", data_point_name="mean", quintile="Total")
        if all(q_rows) and t_rows:
            q_mean = np.mean([wb.current(CV, r[0]) for r in q_rows])
            total = wb.current(CV, t_rows[0])
            if not np.isclose(q_mean, total, rtol=TOTALS_RTOL, atol=0):
                problems.append(f"{c} {v}: amount Total {total:.3g} vs quintile mean {q_mean:.3g} fails the "
                                "prep_extracted totals check")
        # Coverage notebook: intervention target >= baseline coverage
        fort = {wb.get(CV, r, "quintile"): wb.current(CV, r) for r in wb.find(CV, country=c, vehicle=v, need="fortifiability")}
        for f in arm["fortificants"]:
            for row in wb.find(CVF, country=c, vehicle=v, fortificant=f, need="baseline_any"):
                q = wb.get(CVF, row, "quintile")
                fv = fort.get(q, fort.get(ALL_SAME, fort.get("Total")))
                for irow in wb.find(SCEN, country=c, vehicle=v, fortificant=f, need="intervention_coverage"):
                    target = wb.current(SCEN, irow) * fv
                    baseline = wb.current(CVF, row)
                    if target < baseline and not np.isclose(target, baseline):
                        problems.append(f"{c} {v} {f} {q}: intervention target {target:.3g} < baseline {baseline:.3g}")
    return problems


# --------------------------------------------------------------------------------------
# Review report
# --------------------------------------------------------------------------------------


def lit_workbook_rows(wb, r):
    """Workbook rows a literature row speaks to (for comparison only)."""
    if r.need not in NEEDS:
        return []
    sheet = NEEDS[r.need][0]
    criteria = {"need": r.need, "vehicle": r.vehicle, "data_point_name": r.data_point_name}
    if sheet == VEH:
        return wb.find(sheet, **criteria)
    criteria["country"] = r.country
    if sheet in (CVF, SCEN):
        criteria["fortificant"] = r.fortificant
    if sheet == SCEN and norm(r.scenario).startswith("intervention"):
        criteria["scenario"] = r.scenario
    if sheet == CV and not blank(r.sex):
        rows = wb.find(sheet, sex=r.sex, **criteria) or (
            wb.find(sheet, sex=ALL_SAME, **criteria) if r.sex == "Total" else [])
    else:
        rows = wb.find(sheet, **criteria)
    if sheet == SCEN:
        return rows
    q = r.quintile if not blank(r.quintile) else "Total"
    exact = [x for x in rows if norm(wb.get(sheet, x, "quintile")) == norm(q)]
    if exact or q != "Total":
        return exact
    return [x for x in rows if norm(wb.get(sheet, x, "quintile")) == norm(ALL_SAME)]


def recommended_first(lits):
    """'Use in model' = yes candidates first, then unmarked, then 'no'."""
    order = {"yes": 0, "no": 2}
    return sorted(lits, key=lambda x: order.get(x.use_in_model, 1))


def recommendation_tag(x):
    return {"yes": " [recommended]", "no": " [not recommended]"}.get(x.use_in_model, "")


def build_review(wb, lit, arms, decisions, decided, gf_values):
    issues = []
    consumer_basis = among_consumers_arms(decisions)

    def add(kind, detail, **kw):
        issues.append({"issue": kind, "detail": detail, **kw})

    for (sheet, row), entry in wb.plan.items():
        if entry["action"] == "keep" and entry["origin"] == "gf" and entry["method"].startswith("GF national g/cap"):
            add("GF national inconsistent with GF quintiles", entry["method"]
                + ". To use the national figure, add a decision with source method:scale_to_gf_total",
                tab=sheet, row=row, output_value=wb.current(sheet, row))
        # The literature default won over a different GF value, and no decision says which is right
        if entry["origin"] == "lit" and (sheet, row) in gf_values and (sheet, row) not in decided:
            gf_value, gf_ref = gf_values[(sheet, row)]
            if not np.isclose(entry["value"], gf_value, rtol=AGREE_RTOL, atol=1e-9):
                label = " ".join(str(v) for k, v in wb.describe(sheet, row).items() if not blank(v))
                add("GF disagrees with literature default", f"{label}: literature {entry['value']:g} (used) vs "
                    f"GF {gf_value:g} ({gf_ref})", tab=sheet, row=row, output_value=entry["value"],
                    lit_id=entry["ref"].split(";")[0].split(" ")[0])
    for _, r in lit[lit.status == "needs_attention"].iterrows():
        add("literature row needs attention", r.issues, lit_id=r.lit_id,
            location=f"{r.source_sheet} row {r.source_row}")

    by_target = {}
    for _, r in lit[lit.status == "ok"].iterrows():
        rows = lit_workbook_rows(wb, r)
        if not rows:
            if r.need in NEEDS and NEEDS[r.need][0] != VEH:
                add("literature value has no row in the extraction sheet", f"{r.need} {r.country} {r.vehicle} {r.quintile or ''} "
                    f"{r.data_point_name}: {r.value:g}", lit_id=r.lit_id)
            continue
        for row in rows:
            sheet = NEEDS[r.need][0]
            by_target.setdefault((sheet, row), []).append(r)

    for (sheet, row), lits in sorted(by_target.items()):
        current = wb.current(sheet, row)
        current = float(current) if isinstance(current, (int, float)) else current
        entry = wb.plan.get((sheet, row))
        origin = entry["origin"] if entry and entry["action"] == "set" else "existing"
        source = norm(wb.get(sheet, row, "data source"))
        placeholder = origin == "derived" or (origin == "existing" and source.startswith(("dummy", "derived")))
        desc = wb.describe(sheet, row)
        label = " ".join(str(desc[k]) for k in ["country", "vehicle", "fortificant", "scenario", "quintile", "sex",
                                                "data_point_name"] if not blank(desc[k]))
        label = f"{label} | {NEEDS[need_alias(desc['data_need'])][1] if need_alias(desc['data_need']) else ''}"
        lits = recommended_first(lits)
        values = sorted({round(float(x.value), 9) for x in lits})
        recommended = [x for x in lits if x.use_in_model == "yes"]
        if len(values) > 1 and (sheet, row) not in decided:
            add("several literature candidates", f"{label}: " + ", ".join(
                f"{x.lit_id}={x.value:g} ({x.data_source}){recommendation_tag(x)}" for x in lits),
                tab=sheet, row=row, output_value=current, decision_id=decided.get((sheet, row)),
                lit_recommended="; ".join(x.lit_id for x in recommended) or None)
            if len({round(float(x.value), 9) for x in recommended}) > 1:
                add("conflicting literature recommendations", f"{label}: more than one candidate with different "
                    "values is marked 'Use in model' = yes: " + ", ".join(f"{x.lit_id}={x.value:g}" for x in recommended),
                    tab=sheet, row=row, output_value=current)
        for x in lits:
            if x.use_in_model == "no" and recommended:
                continue  # another candidate is marked 'Use in model'; that one is reported instead
            comparable = x.value
            note = ""
            if x.need in ("amount", "u5_amount") and x.data_point_name == "mean" and \
                    (x.country, x.vehicle) in consumer_basis:
                q = wb.get(sheet, row, "quintile") if x.need == "amount" else "Total"
                any_row = wb.find(CV, country=x.country, vehicle=x.vehicle, need="any", quintile=q)
                if any_row:
                    comparable = x.value * wb.current(CV, any_row[0])
                    note = " (literature value x coverage, as this arm's amounts are converted with times_coverage)"
            agree = isinstance(current, (int, float)) and np.isclose(current, comparable, rtol=AGREE_RTOL, atol=1e-9)
            if agree or (sheet, row) in decided:
                continue
            if x.use_in_model == "yes":
                kind = "recommended literature value differs from output"
            elif placeholder:
                kind = "literature could replace placeholder"
            else:
                kind = "literature disagrees with output"
            where = {"existing": "base value", "lit": "literature default", "gf": "GF default",
                     "derived": "derived placeholder", "decision": "decision"}[origin]
            add(kind, f"{label}: output {current!r} ({where}) vs literature {comparable:g}{note} "
                f"[{x.data_source}]{recommendation_tag(x)}", tab=sheet, row=row, output_value=current,
                lit_id=x.lit_id, lit_value=comparable, lit_recommended=x.use_in_model)

    for _, d in decisions[decisions.status.str.lower() == "proposed"].iterrows():
        try:
            effects = resolve_decision(wb, lit, d, gf_values)
        except ValueError as e:
            add("proposed decision is invalid", str(e), decision_id=d.decision_id)
            continue
        for e in effects:
            sheet, row, action, value = e["sheet"], e["row"], e["action"], e["value"]
            now = wb.current(sheet, row)
            now = float(now) if isinstance(now, (int, float)) else now
            add("proposed decision (not applied)", f"{d.rationale} -> would {action} "
                f"{'' if value is None else f'{value:g} '}(output now {now!r})",
                decision_id=d.decision_id, tab=sheet, row=row, output_value=now,
                lit_id=d.source if str(d.source).startswith("lit:") else None, lit_value=value)

    modeled = {(a["country"], a["vehicle"]) for a in arms if a["modeled"]}
    for sheet in (CV, CVF, SCEN, VEH):
        for row in wb.rows(sheet):
            entry = wb.plan.get((sheet, row))
            source = norm(wb.get(sheet, row, "data source"))
            # A placeholder deliberately kept by a decision is no longer reported
            kept_on_purpose = entry is not None and entry["action"] == "keep" and entry["origin"] == "decision"
            is_placeholder = (entry is not None and entry["origin"] == "derived") or (
                (entry is None or entry["action"] == "keep") and source.startswith("dummy")
                and not kept_on_purpose)
            key = (wb.get(sheet, row, "country"), wb.get(sheet, row, "vehicle"))
            if is_placeholder and (sheet == VEH or key in modeled):
                desc = wb.describe(sheet, row)
                add("placeholder remains", " ".join(str(v) for v in desc.values() if not blank(v)),
                    tab=sheet, row=row, output_value=wb.current(sheet, row))
    return pd.DataFrame(issues)


# --------------------------------------------------------------------------------------
# Comparison: base vs GF vs literature (plus the output)
# --------------------------------------------------------------------------------------


def gf_only(gf, base_dir, arms, modeled_only=False):
    """What the GF mapping alone would put in every row.

    By default this covers every arm in the tables, so you can see GF's numbers for arms the
    model doesn't run yet. Rows the mapping can't place are skipped.
    """
    wb = Workbook(base_dir)
    for arm in arms:
        if modeled_only and not arm["modeled"]:
            continue
        arm = dict(arm, modeled=True)
        steps = [] if arm["hces"] else [gf_consumption_any, gf_consumption_amount]
        for step in steps + [gf_fortifiability]:
            try:
                step(gf, wb, arm)
            except (LookupError, TypeError):
                pass
        for f in arm["fortificants"]:
            for step in (gf_baseline, gf_intervention):
                try:
                    step(gf, wb, arm, f)
                except (LookupError, TypeError):
                    pass
    return wb


def compare(a, b):
    if not isinstance(a, (int, float)) or not isinstance(b, (int, float)) or pd.isna(a) or pd.isna(b):
        return ""
    if np.isclose(a, b, rtol=AGREE_RTOL, atol=1e-9):
        return "same"
    return f"differs ({(b - a) / a:+.0%})" if a else "differs"


def build_comparison(gf, wb, lit, arms, decisions, base_dir):
    gfw = gf_only(gf, base_dir, arms)
    consumer_basis = among_consumers_arms(decisions)
    listed = {(a["country"], a["vehicle"]): ("yes" if a["modeled"] else "no") for a in arms}

    lit_by_row, lit_unplaced = {}, []
    for _, r in lit[lit.status.isin(["ok", "needs_attention"])].iterrows():
        rows = lit_workbook_rows(wb, r)
        if not rows:
            if r.need in NEEDS or r.status == "ok":
                lit_unplaced.append(r)
            continue
        for row in rows:
            lit_by_row.setdefault((NEEDS[r.need][0], row), []).append(r)

    records = []
    for sheet in (CV, CVF, SCEN, VEH):
        for row in wb.rows(sheet):
            alias = need_alias(wb.get(sheet, row, "data need"))
            if alias is None:
                continue
            d = wb.describe(sheet, row)
            existing = wb.get(sheet, row, "value")
            g = gfw.plan.get((sheet, row))
            gf_value = g["value"] if g and g["action"] == "set" else None
            lits = recommended_first(lit_by_row.get((sheet, row), []))
            lit_values, lit_published, lit_raw = [], [], []
            for x in lits:
                v = x.value
                lit_raw.append(round(float(x.value), 9))  # in our units, before any basis conversion
                if x.need in ("amount", "u5_amount") and x.data_point_name == "mean" and \
                        (x.country, x.vehicle) in consumer_basis:
                    q = wb.get(sheet, row, "quintile") if x.need == "amount" else "Total"
                    any_row = wb.find(CV, country=x.country, vehicle=x.vehicle, need="any", quintile=q)
                    v = x.value * wb.current(CV, any_row[0]) if any_row else v
                lit_values.append(round(float(v), 9))
                lit_published.append(x.raw_value)
            lit_value = lit_values[0] if len(set(lit_values)) == 1 else None
            rec_values = {v for v, x in zip(lit_values, lits) if x.use_in_model == "yes"}
            rec_ids = [x.lit_id for x in lits if x.use_in_model == "yes"]
            lit_recommended = rec_values.pop() if len(rec_values) == 1 else None
            # Compare against the 'Use in model' pick if there is one, else the single candidate
            lit_compare = lit_recommended if lit_recommended is not None else lit_value
            # GF values are as published, so compare them with the literature as published too
            raw_rec = {v for v, x in zip(lit_raw, lits) if x.use_in_model == "yes"}
            lit_raw_compare = raw_rec.pop() if len(raw_rec) == 1 else (lit_raw[0] if len(set(lit_raw)) == 1 else None)
            entry = wb.plan.get((sheet, row))
            out_set = entry is not None and entry["action"] == "set"
            output = entry["value"] if out_set else existing
            records.append({
                "tab": sheet, "row": row, "need": alias,
                **{k: d[k] for k in ["country", "vehicle", "fortificant", "scenario", "quintile", "sex", "data_point_name"]},
                "units": wb.get(sheet, row, "units"),
                "arm_modeled": listed.get((d["country"], d["vehicle"]), "" if sheet == VEH else "no"),
                "base_value": existing,
                "base_source": wb.get(sheet, row, "data source"),
                "gf_value": gf_value,
                "gf_reference": g["ref"] if g else None,
                "lit_value": lit_value if lit_value is not None else ("; ".join(
                    f"{v:g}{'*' if x.use_in_model == 'yes' else ''}" for v, x in zip(lit_values, lits)) or None),
                "lit_recommended_value": lit_recommended,
                "lit_recommended_id": "; ".join(rec_ids) or None,
                "lit_as_published": "; ".join(str(v) for v in lit_published) or None,
                "lit_ids": "; ".join(x.lit_id for x in lits) or None,
                "lit_source": "; ".join(str(x.data_source) for x in lits) or None,
                "gf_vs_base": compare(existing, gf_value),
                "lit_vs_base": compare(existing, lit_compare),
                "lit_vs_gf": compare(gf_value, lit_raw_compare),
                "output_value": output,
                "output_source": entry["source"] if entry else "base extraction table",
                "decision_id": entry["decision_id"] if entry else None,
            })
    for x in lit_unplaced:
        records.append({"tab": None, "row": None, "need": x.need, "country": x.country, "vehicle": x.vehicle,
                        "fortificant": x.fortificant, "scenario": x.scenario, "quintile": x.quintile, "sex": x.sex,
                        "data_point_name": x.data_point_name, "units": x.units,
                        "arm_modeled": listed.get((x.country, x.vehicle), "no"),
                        "lit_value": x.value, "lit_as_published": x.raw_value, "lit_ids": x.lit_id,
                        "lit_source": x.data_source,
                        "output_source": "(no row in the extraction sheet)" if x.status == "ok"
                        else f"(literature row needs attention: {x.issues})"})
    return pd.DataFrame(records)


def write_comparison_xlsx(df, path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    df.to_excel(path, index=False, sheet_name="comparison")
    book = openpyxl.load_workbook(path)
    ws = book.active
    ws.freeze_panes = "L2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True)
    cols = {c.value: c.column for c in ws[1]}
    differs, same = PatternFill("solid", fgColor="F8CBAD"), PatternFill("solid", fgColor="C6EFCE")
    for r in range(2, ws.max_row + 1):
        for flag, target in [("gf_vs_base", "gf_value"), ("lit_vs_base", "lit_value")]:
            v = ws.cell(r, cols[flag]).value or ""
            fill = differs if v.startswith("differs") else same if v == "same" else None
            if fill:
                ws.cell(r, cols[target]).fill = fill
                ws.cell(r, cols[flag]).fill = fill
        v = ws.cell(r, cols["lit_vs_gf"]).value or ""
        if v:
            ws.cell(r, cols["lit_vs_gf"]).fill = differs if v.startswith("differs") else same
    widths = {"need": 22, "base_source": 30, "gf_reference": 30, "lit_source": 30, "lit_ids": 16,
              "output_source": 24}
    for name, col in cols.items():
        ws.column_dimensions[get_column_letter(col)].width = widths.get(name, 12)
    book.save(path)


# --------------------------------------------------------------------------------------


def read_decisions(path):
    d = pd.read_csv(path, dtype=str, encoding="utf-8-sig").fillna("")
    d.columns = [c.strip() for c in d.columns]
    if d.decision_id.duplicated().any():
        raise ValueError(f"duplicate decision_id: {d.decision_id[d.decision_id.duplicated()].tolist()}")
    return d


class Build:
    """Everything one build produces (see build())."""


def build(base_dir=BASE_DIR, lit_long=LIT_LONG, gf_long=GF_LONG, model_arms_file=MODEL_ARMS_FILE,
          decisions_file=DECISIONS_FILE):
    """Run phases 1 to 4 and the reports. Writes nothing."""
    global GF_DATA
    b = Build()
    gf, lit = GF(gf_long), pd.read_csv(lit_long)
    GF_DATA = gf
    if "use_in_model" not in lit.columns:
        lit["use_in_model"] = None
    decisions = read_decisions(decisions_file)
    wb = Workbook(base_dir)
    arms = read_arms(wb, model_arms_file)

    # Phase 1: defaults. gf_values holds every GF default, used or not, so that 'gf'
    # decisions can choose GF over the literature.
    gfw = gf_only(gf, base_dir, arms, modeled_only=True)
    gf_values = {key: (e["value"], e["ref"]) for key, e in gfw.plan.items() if e["action"] == "set"}
    # The literature first (except for 2035 targets that GF gives)...
    ambiguous = apply_lit_defaults(wb, lit, arms, gf_values)
    # ...then GF for every row the literature left unset
    gf_automatic = []  # places where a GF rule deliberately did NOT use GF's number
    for key, e in sorted(gfw.plan.items()):
        if key in wb.plan:
            continue
        wb.plan[key] = e
        if e["action"] == "keep":
            gf_automatic.append({"tab": key[0], "row": key[1], "reason": e["method"], "ref": e["ref"]})
    # Phases 2 and 3: decisions
    decided = apply_decisions(wb, lit, decisions, gf_values)

    b.wb, b.lit, b.arms, b.decisions = wb, lit, arms, decisions
    b.gf_values, b.gf_automatic = gf_values, gf_automatic
    b.review = enrich_review(build_review(wb, lit, arms, decisions, decided, gf_values), wb, lit, decisions)
    b.comparison = build_comparison(gf, wb, lit, arms, decisions, base_dir)
    # Phase 4: checks. Disagreeing literature candidates that no decision resolves fail the
    # build, so a row never silently switches source when a candidate is added.
    unresolved = {key: c for key, c in ambiguous.items() if key not in decided}
    b.problems = run_checks(wb, arms, unresolved)
    records, frames = wb.outputs(SOURCE_LABELS)
    b.changelog = pd.DataFrame(records)
    b.tables = {t: table_text(df) for t, df in frames.items()}
    return b


def value_changes(tables, generated_dir=GENERATED_DIR):
    """Rows whose Value differs from the generated tables on disk, as readable lines."""
    from common import read_table
    out = []
    for t, text in tables.items():
        path = pathlib.Path(generated_dir) / f"{t}.csv"
        if not path.exists():
            continue
        old = read_table(path)
        new = read_table(io.StringIO(text))
        if len(old) != len(new):
            out.append(f"{t}: {len(old)} -> {len(new)} rows")
            continue
        keys = [c for c in ["Country", "Vehicle", "Fortificant", "Scenario", "Quintile", "Sex", "Data need",
                            "Data point name"] if c in new.columns]
        for i in range(len(new)):
            if old.at[i, "Value"] != new.at[i, "Value"]:
                label = " ".join(new.at[i, k] for k in keys if new.at[i, k])
                out.append(f"{t} row {i + 2} ({label}): {old.at[i, 'Value'] or '(blank)'} -> "
                           f"{new.at[i, 'Value'] or '(blank)'}")
    return out


def stale_tables(tables, generated_dir=GENERATED_DIR):
    """Generated tables whose committed file differs from a fresh build."""
    stale = []
    for t, text in tables.items():
        path = pathlib.Path(generated_dir) / f"{t}.csv"
        if not path.exists() or path.read_text(encoding="utf-8") != text:
            stale.append(path.name)
    return stale


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--base", type=pathlib.Path, default=BASE_DIR, help="directory of the base tables")
    parser.add_argument("--output-dir", type=pathlib.Path, default=GENERATED_DIR,
                        help="where to write the generated tables")
    parser.add_argument("--decisions", type=pathlib.Path, default=DECISIONS_FILE)
    parser.add_argument("--model-arms", type=pathlib.Path, default=MODEL_ARMS_FILE,
                        help="the arms the model runs (default: 0050_config/location_fortificant_vehicles.csv)")
    parser.add_argument("--lit-long", type=pathlib.Path, default=LIT_LONG)
    parser.add_argument("--check", action="store_true",
                        help="write nothing; exit 1 if the generated tables are out of date")
    args = parser.parse_args()
    if args.output_dir.resolve() == args.base.resolve():
        sys.exit("Refusing to overwrite the base tables; choose another --output-dir")

    b = build(args.base, args.lit_long, GF_LONG, args.model_arms, args.decisions)
    if b.problems:
        print("Checks failed; tables NOT written:")
        for p in b.problems:
            print("  " + p)
        if not args.check:
            b.review.to_csv(REVIEW, index=False)
            print(f"See {REVIEW.name} for context.")
        sys.exit(1)
    if args.check:
        stale = stale_tables(b.tables, args.output_dir)
        if stale:
            sys.exit(f"Out of date in {args.output_dir}: {', '.join(stale)}. Run build_extraction.py.")
        print("Generated tables are up to date.")
        return

    changes = value_changes(b.tables, args.output_dir)
    if changes:
        print(f"VALUES CHANGED in {args.output_dir.name}/ ({len(changes)}):")
        for c in changes:
            print("  " + c)
    else:
        print(f"No values changed in {args.output_dir.name}/.")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    for t, text in b.tables.items():
        (args.output_dir / f"{t}.csv").write_text(text, encoding="utf-8")
    b.changelog.to_csv(CHANGELOG, index=False, lineterminator="\n")
    b.review.to_csv(REVIEW, index=False, lineterminator="\n")
    write_status(STATUS, b.wb, b.lit, b.arms, b.decisions, b.review, resolve_decision, b.gf_values, b.gf_automatic)
    b.comparison.to_csv(COMPARISON, index=False, lineterminator="\n")
    written = [f"{args.output_dir.name}/*.csv", CHANGELOG.name, REVIEW.name, STATUS.name, COMPARISON.name]
    try:
        write_comparison_xlsx(b.comparison, COMPARISON_XLSX)
        written.append(COMPARISON_XLSX.name)
    except ImportError:
        print("openpyxl not installed; skipped comparison.xlsx")

    print(f"Wrote {', '.join(written)}")
    print(b.changelog.groupby(["step", "source", "status"]).size().to_string())
    if len(b.review):
        print("\nReview items:")
        print(b.review.groupby("issue").size().to_string())


if __name__ == "__main__":
    main()
