"""Bring Gates Foundation background data into the extraction workbook.

Reads ``Nutrition PST Background Data.xlsx`` (the GF "LSFF background data" file) and
writes an updated copy of ``Data Extraction Sheet.xlsx`` along with a change log
listing every cell it touched (old value, new value, GF source cell, method).

The original extraction workbook is never modified. Review the change log, then
replace the original with the new copy yourself.

Usage (from this directory):

    python import_gf_background_data.py
    python import_gf_background_data.py --output "Data Extraction Sheet (GF).xlsx"

Terminology (GF -> ours). See GF_DATA_MAPPING.md for the full explanation.

    GF "Coverage" (Cov)         -> "Vehicle consumption by WRA -- any"
    GF "g/cap"                  -> "Vehicle consumption by WRA -- amount" (mean)
                                   (x Coverage for vehicles where the NFCMS figure
                                   is among consumers: rice, wheat)
    GF "Consolidation" (2035)   -> 'Vehicle "fortifiability" ...'
    GF current Cons x Compl     -> baseline "any" coverage x "Baseline effective %"
    GF "Compliance" (2035)      -> "Intervention coverage % of fortifiable" x
                                   "Intervention effective % of fortified"
                                   (equal split: each = sqrt(compliance), 2 d.p.)

Formula cells: openpyxl cannot compute formulas, and a workbook it saves has no
cached formula results, which ``pandas.read_excel`` (used by the pipeline) needs.
So after saving, the script recalculates the copy with LibreOffice (``soffice``) and
then checks that every cell it did not mean to change still has its old value.
If LibreOffice is not installed, open the new file in Excel and save it before
using it.
"""

import argparse
import math
import pathlib
import shutil
import subprocess
import sys
import tempfile

import numpy as np
import openpyxl
import pandas as pd

HERE = pathlib.Path(__file__).resolve().parent
GF_FILE = HERE / "Nutrition PST Background Data.xlsx"
EXTRACTION_FILE = HERE / "Data Extraction Sheet.xlsx"
DEFAULT_OUTPUT = HERE / "Data Extraction Sheet (GF 2026-09-30).xlsx"
DEFAULT_CHANGELOG = HERE / "gf_import_changelog.csv"

SOURCE_LABEL = "GF Nutrition PST Background Data (updated 30 Sep 2026)"
TARGET_YEAR = 2035

# --------------------------------------------------------------------------------------
# Terminology tables
# --------------------------------------------------------------------------------------

# GF vehicle labels (they differ slightly between tabs) -> our vehicle names
GF_VEHICLES = {
    "wheat flour": "Wheat",
    "rice": "Rice",
    "salt": "Salt",
    "salt (dfs)": "Salt",
    "salt/dfs": "Salt",
    "bouillon": "Bouillon",
    "edible oil": "Oil",
    "maize flour": "Maize",
}

# GF wealth-quintile labels (they differ between country tabs) -> our labels
GF_QUINTILES = {
    "poorest": "Lowest",
    "lowest": "Lowest",
    "second": "Second",
    "2nd": "Second",
    "third": "Middle",
    "3rd": "Middle",
    "fourth": "Fourth",
    "4th": "Fourth",
    "wealthiest": "Highest",
}
QUINTILES = ["Lowest", "Second", "Middle", "Fourth", "Highest"]

# The arms we model (must match 0050_config/location_fortificant_vehicles.csv), and how
# to interpret GF's g/cap for each. "among_consumers": NFCMS reports intake among
# consumers, so we multiply by coverage to get the population mean the model expects
# (JG email 8/31 for rice). "per_capita": already a mean over all women (NFCMS bouillon
# tables; Ethiopia salt with ~98% coverage).
ARMS = {
    ("Nigeria", "Rice"): {"fortificants": ["Iron", "Folate"], "gcap_basis": "among_consumers"},
    ("Nigeria", "Bouillon"): {"fortificants": ["Iron", "Folate"], "gcap_basis": "per_capita"},
    ("Nigeria", "Wheat"): {"fortificants": ["Iron", "Folate"], "gcap_basis": "among_consumers"},
    ("Ethiopia", "Salt"): {"fortificants": ["Folate"], "gcap_basis": "per_capita"},
    # India rice consumption and baseline coverage come from HCES microdata
    # (0100_data_prep/hces), not from the workbook, so only the intervention rows apply.
    ("India", "Rice"): {"fortificants": ["Iron", "Folate"], "gcap_basis": None},
}

NEED_ANY = "Vehicle consumption by WRA -- any"
NEED_AMOUNT = "Vehicle consumption by WRA -- amount"
NEED_U5_ANY = "Vehicle consumption by U5 children -- any"
NEED_U5_AMOUNT = "Vehicle consumption by U5 children -- amount"
NEED_FORTIFIABILITY = 'Vehicle "fortifiability" (essentially amount industrially produced)'
NEED_BASE_ANY = "Vehicle fortification at baseline -- any"
NEED_BASE_EFF = "Baseline effective % of fortified"
NEED_INT_COV = "Intervention coverage % of fortifiable"
NEED_INT_EFF = "Intervention effective % of fortified"

CV_SHEET = "Country-Vehicle Extraction"
CVF_SHEET = "Country-Vehicle-Fort Extraction"
SCEN_SHEET = "Scenario Definition Extraction"

# Same tolerance as check_totals_reasonable in prep_extracted.ipynb
TOTALS_RTOL = 0.1


def norm(x):
    return str(x).strip().lower() if x is not None else ""


def round2(x):
    """Round half up to 2 d.p., the precision of the existing intervention rows."""
    return math.floor(x * 100 + 0.5) / 100


def as_number(v):
    """GF uses 'n/a' (and occasionally '<1%') for missing values."""
    if v is None or isinstance(v, str):
        return None
    return float(v)


# --------------------------------------------------------------------------------------
# Reading the GF workbook
# --------------------------------------------------------------------------------------


class GFData:
    """Lookups into the GF workbook that only ever read typed-in cells.

    The GF file links many cells by formula (e.g. country tabs' National rows, the
    "Current" columns of Coverage Over Time), and it was last saved with
    fullCalcOnLoad set, so cached formula results may be stale. We therefore read the
    cells those formulas point at, and refuse to read any formula cell.
    """

    def __init__(self, path):
        self.path = path
        self.wb = openpyxl.load_workbook(path, data_only=False)

    def _value(self, ws, row, col):
        cell = ws.cell(row, col)
        if isinstance(cell.value, str) and cell.value.startswith("="):
            raise ValueError(
                f"{ws.title}!{cell.coordinate} is a formula ({cell.value}); "
                "refusing to rely on a possibly stale cached value"
            )
        return cell.value, f"'{ws.title}'!{cell.coordinate}"

    @staticmethod
    def _header_row(ws):
        """The row whose first cell is 'Country' / 'Stratum'; vehicle names sit just above."""
        return next(
            r for r in range(1, 30) if norm(ws.cell(r, 1).value) in ("country", "stratum")
        )

    @classmethod
    def _vehicle_columns(cls, ws):
        """Map (vehicle, sub-header) -> column, for the two-row vehicle headers."""
        sub_row = cls._header_row(ws)
        group_row = sub_row - 1
        columns = {}
        group = None
        for col in range(1, ws.max_column + 1):
            if ws.cell(group_row, col).value is not None:
                group = norm(ws.cell(group_row, col).value)
            sub = norm(ws.cell(sub_row, col).value)
            if group in GF_VEHICLES and sub:
                columns[(GF_VEHICLES[group], sub)] = col
        return columns

    def national(self, country, vehicle, metric):
        """Current national value from 'Burden x Vehicle Matrix'.

        metric is one of 'cons', 'cov', 'g/cap', 'compl'.
        """
        ws = self.wb["Burden x Vehicle Matrix"]
        cols = self._vehicle_columns(ws)
        for row in range(self._header_row(ws) + 1, ws.max_row + 1):
            if norm(ws.cell(row, 1).value) == country.lower():
                value, ref = self._value(ws, row, cols[(vehicle, metric)])
                return as_number(value), ref, value
        raise KeyError(country)

    def by_quintile(self, country, vehicle, metric):
        """Wealth-quintile values from the '<Country> (stratified)' tab."""
        ws = self.wb[f"{country} (stratified)"]
        cols = self._vehicle_columns(ws)
        col = cols[(vehicle, metric)]
        start = next(
            r for r in range(1, ws.max_row + 1) if norm(ws.cell(r, 1).value) == "wealth quintile"
        )
        out = {}
        for row in range(start + 1, start + 6):
            quintile = GF_QUINTILES[norm(ws.cell(row, 1).value)]
            value, ref = self._value(ws, row, col)
            out[quintile] = (as_number(value), ref)
        assert list(out) == QUINTILES, out.keys()
        return out

    def target(self, country, vehicle, metric, year=TARGET_YEAR):
        """Typed-in 2031/2035 value from 'Coverage Over Time'.

        metric is one of 'coverage', 'consolidation', 'compliance'.
        """
        ws = self.wb["Coverage Over Time"]
        header_row = next(
            r for r in range(1, 20) if norm(ws.cell(r, 1).value) == "country"
        )
        col = next(
            c
            for c in range(1, ws.max_column + 1)
            if norm(ws.cell(header_row, c).value).replace("\n", " ") == f"{year} {metric}"
        )
        for row in range(header_row + 1, ws.max_row + 1):
            if (
                norm(ws.cell(row, 1).value) == country.lower()
                and GF_VEHICLES.get(norm(ws.cell(row, 2).value)) == vehicle
            ):
                value, ref = self._value(ws, row, col)
                return as_number(value), ref, value
        raise KeyError((country, vehicle))


# --------------------------------------------------------------------------------------
# Editing the extraction workbook
# --------------------------------------------------------------------------------------


class Extraction:
    def __init__(self, path):
        self.path = path
        self.wb = openpyxl.load_workbook(path)  # formulas, for writing
        self.wb_values = openpyxl.load_workbook(path, data_only=True)  # cached values
        self.changes = []
        self.expected = {}  # (sheet, coordinate) -> expected value after recalc

    def header(self, sheet):
        ws = self.wb_values[sheet]
        return {norm(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}

    def find(self, sheet, **criteria):
        """Row numbers whose columns match all criteria (case/space-insensitive)."""
        ws = self.wb_values[sheet]
        header = self.header(sheet)
        rows = []
        for row in range(2, ws.max_row + 1):
            if all(
                norm(ws.cell(row, header[norm(col.replace("_", " "))]).value) == norm(val)
                for col, val in criteria.items()
            ):
                rows.append(row)
        return rows

    def get(self, sheet, row, column):
        return self.wb_values[sheet].cell(row, self.header(sheet)[norm(column)]).value

    def describe(self, sheet, row):
        cols = ["country", "vehicle", "fortificant", "scenario", "quintile", "sex", "data need", "data point name"]
        header = self.header(sheet)
        return {c.replace(" ", "_"): self.get(sheet, row, c) if c in header else None for c in cols}

    def set(self, sheet, row, value, *, gf_ref, method, formula=None, used=True):
        """Set the Value cell (and provenance columns) of one row.

        value is the numeric result; formula, if given, is written instead so the
        workbook keeps showing how the number was derived.
        """
        header = self.header(sheet)
        ws = self.wb[sheet]
        value_col = header["value"]
        old = self.get(sheet, row, "value")
        old_source = self.get(sheet, row, "data source")
        rel = None
        if isinstance(old, (int, float)):
            rel = (value - old) / old if old else (0.0 if value == 0 else np.inf)
        unchanged = isinstance(old, (int, float)) and np.isclose(old, value, rtol=1e-9, atol=1e-12)

        record = {
            "sheet": sheet,
            "row": row,
            **self.describe(sheet, row),
            "old_value": old,
            "new_value": value,
            "relative_change": rel,
            "old_source": old_source,
            "gf_reference": gf_ref,
            "method": method,
            "used_by_model": used,
            "status": "unchanged" if unchanged else "changed",
        }
        self.changes.append(record)
        if unchanged:
            return

        ws.cell(row, value_col).value = formula if formula is not None else value
        self.expected[(sheet, ws.cell(row, value_col).coordinate)] = value
        provenance = {
            # Derived placeholders stay easy to find by filtering on Data source
            "data source": "DERIVED (MIC-7549)" if method.startswith("DERIVED") else SOURCE_LABEL,
            "notes": f"{gf_ref}; {method} (was {old!r} from {old_source!r})",
        }
        # A CI or SE from the previous source no longer describes a materially new value
        if rel is None or abs(rel) > 0.01:
            provenance.update({"ci": None, "se": None})
        for col, new in provenance.items():
            if col in header:
                ws.cell(row, header[col]).value = new
                self.expected[(sheet, ws.cell(row, header[col]).coordinate)] = new

    def note(self, sheet, rows, *, status, gf_ref, method, used):
        """Record a row we deliberately did not change."""
        for row in rows:
            self.changes.append(
                {
                    "sheet": sheet,
                    "row": row,
                    **self.describe(sheet, row),
                    "old_value": self.get(sheet, row, "value"),
                    "new_value": None,
                    "relative_change": None,
                    "old_source": self.get(sheet, row, "data source"),
                    "gf_reference": gf_ref,
                    "method": method,
                    "used_by_model": used,
                    "status": status,
                }
            )

    def one(self, sheet, **criteria):
        rows = self.find(sheet, **criteria)
        if len(rows) != 1:
            raise LookupError(f"Expected 1 row in {sheet} for {criteria}, found {rows}")
        return rows[0]


# --------------------------------------------------------------------------------------
# The mapping, one data need at a time
# --------------------------------------------------------------------------------------


def update_consumption_any(gf, ex, country, vehicle):
    """GF Coverage -> % of WRA consuming the vehicle, by quintile plus Total."""
    national, nat_ref, _ = gf.national(country, vehicle, "cov")
    by_q = gf.by_quintile(country, vehicle, "cov")
    have_quintiles = all(v is not None for v, _ in by_q.values())
    rows_by_quintile = {}

    for need in [NEED_ANY] + ([NEED_U5_ANY] if vehicle == "Wheat" else []):
        # U5 'any' is never read by the pipeline; we only refresh the wheat rows, which
        # were placeholders copying the WRA rows (as rice's U5 rows do).
        used = need == NEED_ANY
        method = "GF Coverage (% using the vehicle)" + ("" if used else "; U5 row mirrors WRA, as for rice")
        if have_quintiles:
            for q, (value, ref) in by_q.items():
                row = ex.one(CV_SHEET, country=country, vehicle=vehicle, data_need=need, quintile=q)
                ex.set(CV_SHEET, row, value, gf_ref=ref, method=method, used=used)
                if used:
                    rows_by_quintile[q] = (row, value)
            row = ex.one(CV_SHEET, country=country, vehicle=vehicle, data_need=need, quintile="Total")
            ex.set(CV_SHEET, row, national, gf_ref=nat_ref,
                   method=method + "; Total is only used to sanity-check the quintiles", used=used)
            if used:
                rows_by_quintile["Total"] = (row, national)
        else:
            # e.g. Ethiopia salt: GF has no quintile breakdown, so keep "All (assumed same)"
            row = ex.one(CV_SHEET, country=country, vehicle=vehicle, data_need=need,
                         quintile="All (assumed same)")
            ex.set(CV_SHEET, row, national, gf_ref=nat_ref,
                   method=method + "; national value, no quintile breakdown in GF", used=used)
    return rows_by_quintile


def update_consumption_amount(gf, ex, country, vehicle, basis, any_rows):
    """GF g/cap -> mean g/day over all WRA (the model's mixture mean, including zeros)."""
    value_letter = openpyxl.utils.get_column_letter(ex.header(CV_SHEET)["value"])
    national, nat_ref, _ = gf.national(country, vehicle, "g/cap")
    by_q = gf.by_quintile(country, vehicle, "g/cap")
    new_means = {}

    for q, (gcap, ref) in by_q.items():
        row = ex.one(CV_SHEET, country=country, vehicle=vehicle, data_need=NEED_AMOUNT,
                     data_point_name="mean", quintile=q)
        if basis == "among_consumers":
            any_row, cov = any_rows[q]
            value = gcap * cov
            formula = f"={gcap}*{value_letter}{any_row}"
            method = "GF g/cap is among consumers (NFCMS), multiplied by WRA coverage"
        else:
            value, formula = gcap, None
            method = "GF g/cap used as the mean over all WRA"
        ex.set(CV_SHEET, row, value, gf_ref=ref, method=method, formula=formula)
        new_means[q] = value

    # The Total row only feeds check_totals_reasonable. GF's national g/cap comes from a
    # different source (M4N) than its strata (NFCMS/IHME), so only use it if it passes
    # the same check the pipeline applies.
    row = ex.one(CV_SHEET, country=country, vehicle=vehicle, data_need=NEED_AMOUNT,
                 data_point_name="mean", quintile="Total")
    if basis == "among_consumers":
        any_row, cov = any_rows["Total"]
        total = national * cov
        formula = f"={national}*{value_letter}{any_row}"
    else:
        total, formula = national, None
    quintile_mean = np.mean(list(new_means.values()))
    if np.isclose(quintile_mean, total, rtol=TOTALS_RTOL, atol=0):
        ex.set(CV_SHEET, row, total, gf_ref=nat_ref, formula=formula,
               method="GF national g/cap (converted as for the quintiles); only used to sanity-check them")
    else:
        ex.note(CV_SHEET, [row], status="kept: GF national inconsistent with GF quintiles",
                gf_ref=nat_ref, used=False,
                method=f"GF national converts to {total:.3g} g/day but the unweighted mean of the "
                       f"quintiles is {quintile_mean:.3g}; the pipeline's totals check (rtol "
                       f"{TOTALS_RTOL}) would fail, so the existing Total was kept")
    return new_means


def consumer_moments(p, mean, sd):
    """Invert the pregnancy sim's zero-inflated normal (intervention.py).

    The workbook's mean/SD describe all women, a mixture of a point mass at 0
    (probability 1 - p) and a normal over consumers. Returns the consumers' mean and
    variance; the variance must be positive or the sim draws NaNs.
    """
    mean_a = mean / p
    var_a = sd**2 / p - mean_a**2 * (1 - p)
    return mean_a, var_a


def mixture_sd(p, mean, consumer_cv):
    """SD over all women given the consumers' CV (inverse of consumer_moments)."""
    mean_a = mean / p
    var_a = (consumer_cv * mean_a) ** 2
    return math.sqrt(p * var_a + mean_a**2 * p * (1 - p))


def update_wheat_sds_and_u5(ex, wheat_means, rice_means):
    """No SDs or U5 amounts in the GF file: derive the wheat placeholders from rice.

    WRA SD: assume wheat consumers have the same coefficient of variation as rice
    consumers in that quintile, then convert back to an SD over all women using wheat
    coverage. (Scaling the all-women SD directly would ignore that wheat coverage is
    much lower, and gives negative consumer variances.)
    U5 mean: rice U5 mean x (wheat/rice ratio of the WRA quintile means).
    U5 SD: as for WRA, with national coverage and the rice U5 consumer CV.
    These remain assumptions; they are tagged "DERIVED (MIC-7549)" in Data source.
    """
    def current(vehicle, need, point, **criteria):
        row = ex.one(CV_SHEET, country="Nigeria", vehicle=vehicle, data_need=need,
                     data_point_name=point, **criteria)
        return new_value(ex, CV_SHEET, row)

    def coverage(vehicle, q):
        return current(vehicle, NEED_ANY, "percentage", quintile=q)

    for q in QUINTILES + ["Total"]:
        p_rice, p_wheat = coverage("Rice", q), coverage("Wheat", q)
        rice_mean_a, rice_var_a = consumer_moments(
            p_rice, current("Rice", NEED_AMOUNT, "mean", quintile=q),
            current("Rice", NEED_AMOUNT, "standard deviation", quintile=q),
        )
        cv = math.sqrt(rice_var_a) / rice_mean_a
        sd = mixture_sd(p_wheat, current("Wheat", NEED_AMOUNT, "mean", quintile=q), cv)
        row = ex.one(CV_SHEET, country="Nigeria", vehicle="Wheat", data_need=NEED_AMOUNT,
                     data_point_name="standard deviation", quintile=q)
        ex.set(CV_SHEET, row, sd, gf_ref="(none: GF has no SDs)",
               method=f"DERIVED: wheat consumers assumed to have the Nigeria rice consumer CV ({cv:.3f}) "
                      f"for this quintile; converted to an SD over all women with wheat coverage {p_wheat}")

    ratio = np.mean([wheat_means[q] for q in QUINTILES]) / np.mean([rice_means[q] for q in QUINTILES])
    p_rice, p_wheat = coverage("Rice", "Total"), coverage("Wheat", "Total")
    for sex in ["Total", "Female", "Male"]:
        rice_mean = current("Rice", NEED_U5_AMOUNT, "mean", sex=sex)
        rice_sd = current("Rice", NEED_U5_AMOUNT, "standard deviation", sex=sex)
        wheat_mean = rice_mean * ratio
        row = ex.one(CV_SHEET, country="Nigeria", vehicle="Wheat", data_need=NEED_U5_AMOUNT,
                     data_point_name="mean", sex=sex)
        ex.set(CV_SHEET, row, wheat_mean, gf_ref="(none: GF has no U5 data)",
               method=f"DERIVED: Nigeria rice U5 mean x {ratio:.4f} (wheat/rice ratio of the mean WRA "
                      "quintile means); only affects ages 5-15 via prep_extracted's interpolation")
        rice_mean_a, rice_var_a = consumer_moments(p_rice, rice_mean, rice_sd)
        cv = math.sqrt(rice_var_a) / rice_mean_a
        row = ex.one(CV_SHEET, country="Nigeria", vehicle="Wheat", data_need=NEED_U5_AMOUNT,
                     data_point_name="standard deviation", sex=sex)
        ex.set(CV_SHEET, row, mixture_sd(p_wheat, wheat_mean, cv), gf_ref="(none: GF has no U5 data)",
               method=f"DERIVED: rice U5 consumer CV ({cv:.3f}) at national coverage; converted with "
                      f"wheat national coverage {p_wheat}")


def check_consumer_variance(ex):
    """The pregnancy sim needs a positive consumer variance in every WRA quintile."""
    problems = []
    for (country, vehicle), spec in ARMS.items():
        if spec["gcap_basis"] is None:
            continue
        for row in ex.find(CV_SHEET, country=country, vehicle=vehicle, data_need=NEED_AMOUNT,
                           data_point_name="mean"):
            q = ex.get(CV_SHEET, row, "quintile")
            sd_rows = ex.find(CV_SHEET, country=country, vehicle=vehicle, data_need=NEED_AMOUNT,
                              data_point_name="standard deviation", quintile=q)
            sd_rows = sd_rows or ex.find(CV_SHEET, country=country, vehicle=vehicle, data_need=NEED_AMOUNT,
                                         data_point_name="standard deviation", quintile="All (assumed same)")
            any_rows = ex.find(CV_SHEET, country=country, vehicle=vehicle, data_need=NEED_ANY, quintile=q)
            any_rows = any_rows or ex.find(CV_SHEET, country=country, vehicle=vehicle, data_need=NEED_ANY,
                                           quintile="All (assumed same)")
            _, var_a = consumer_moments(new_value(ex, CV_SHEET, any_rows[0]), new_value(ex, CV_SHEET, row),
                                        new_value(ex, CV_SHEET, sd_rows[0]))
            if var_a <= 0:
                problems.append((country, vehicle, q, var_a))
    if problems:
        raise AssertionError(f"Non-positive consumer variance (sim would draw NaN): {problems}")


def update_fortifiability(gf, ex, country, vehicle):
    """GF 2035 consolidation -> fortifiability (used only for intervention targets)."""
    value, ref, raw = gf.target(country, vehicle, "consolidation")
    rows = ex.find(CV_SHEET, country=country, vehicle=vehicle, data_need=NEED_FORTIFIABILITY)
    assert rows, (country, vehicle)
    if country == "India":
        ex.note(CV_SHEET, rows, status="skipped: needs decision", gf_ref=ref, used=True,
                method=f"GF {TARGET_YEAR} consolidation is {raw}. In our model India's value is applied "
                       "only to purchased non-PDS rice (prep_extracted 'India industry consolidation'), "
                       "while GF's is market-wide, so it is not a like-for-like replacement")
        return
    if value is None:
        ex.note(CV_SHEET, rows, status="kept: GF n/a", gf_ref=ref, used=True, method="")
        return
    for row in rows:  # Total and any quintile rows all get the national target
        ex.set(CV_SHEET, row, value, gf_ref=ref,
               method=f"GF {TARGET_YEAR} consolidation; national, applied to every quintile")


def update_baseline(gf, ex, country, vehicle, fortificant):
    """GF current consolidation x compliance -> baseline any coverage x effectiveness."""
    any_rows = ex.find(CVF_SHEET, country=country, vehicle=vehicle, fortificant=fortificant,
                       data_need=NEED_BASE_ANY)
    eff_rows = ex.find(CVF_SHEET, country=country, vehicle=vehicle, fortificant=fortificant,
                       data_need=NEED_BASE_EFF)
    cons, cons_ref, cons_raw = gf.national(country, vehicle, "cons")
    compl, compl_ref, compl_raw = gf.national(country, vehicle, "compl")
    ref = f"{cons_ref} x {compl_ref}"

    if country == "India":
        ex.note(CVF_SHEET, eff_rows, status="skipped: needs decision", gf_ref=ref, used=True,
                method=f"GF current consolidation {cons_raw}, compliance {compl_raw}. India baseline "
                       "coverage comes from HCES (GOVERNMENT_BASELINE_COVERAGE = 0.8 in 01_extract_hces) "
                       "so GF's market-wide numbers cannot be mapped onto this row alone")
        return
    if cons is None or compl is None:
        ex.note(CVF_SHEET, any_rows + eff_rows, status="kept: GF n/a", gf_ref=ref, used=True,
                method=f"GF current consolidation {cons_raw!r}, compliance {compl_raw!r}")
        return

    split = round2(math.sqrt(compl))
    coverage = cons * split
    assert len(any_rows) == 1 and len(eff_rows) == 1, (any_rows, eff_rows)
    ex.set(CVF_SHEET, any_rows[0], coverage, gf_ref=ref,
           method=f"consolidation x sqrt(compliance) = {cons} x {split} (equal split of current compliance)")
    if coverage == 0:
        ex.note(CVF_SHEET, eff_rows, status="kept: irrelevant (no baseline coverage)", gf_ref=ref,
                used=True, method="baseline coverage is 0, so baseline effectiveness has no effect")
    else:
        ex.set(CVF_SHEET, eff_rows[0], split, gf_ref=compl_ref,
               method=f"sqrt(current compliance {compl}) rounded to 2 d.p.")


def update_intervention(gf, ex, country, vehicle, fortificant):
    """GF 2035 compliance -> intervention coverage of fortifiable = effectiveness = sqrt."""
    compl, ref, raw = gf.target(country, vehicle, "compliance")
    if compl is None:
        raise ValueError(f"No {TARGET_YEAR} compliance for {country} {vehicle}: {raw!r}")
    split = round2(math.sqrt(compl))
    for need in [NEED_INT_COV, NEED_INT_EFF]:
        rows = ex.find(SCEN_SHEET, country=country, vehicle=vehicle, fortificant=fortificant, data_need=need)
        assert rows, (country, vehicle, fortificant, need)
        for row in rows:  # one per scenario (Ethiopia has three)
            ex.set(SCEN_SHEET, row, split, gf_ref=ref,
                   method=f"sqrt({TARGET_YEAR} compliance {compl}) rounded to 2 d.p. (equal split, as KB 9/9)")


def check_intervention_exceeds_baseline(ex):
    """Mirror the assertion in calculate_effective_coverage_by_quintile_and_scenario."""
    problems = []
    for (country, vehicle), spec in ARMS.items():
        if country == "India":
            continue  # baseline from HCES; checked by the coverage notebook itself
        fort_rows = ex.find(CV_SHEET, country=country, vehicle=vehicle, data_need=NEED_FORTIFIABILITY)
        fortifiability = {ex.get(CV_SHEET, r, "quintile"): new_value(ex, CV_SHEET, r) for r in fort_rows}
        for fortificant in spec["fortificants"]:
            for row in ex.find(CVF_SHEET, country=country, vehicle=vehicle, fortificant=fortificant,
                               data_need=NEED_BASE_ANY):
                q = ex.get(CVF_SHEET, row, "quintile")
                baseline = new_value(ex, CVF_SHEET, row)
                fort = fortifiability.get(q, fortifiability.get("All (assumed same)", fortifiability.get("Total")))
                for int_row in ex.find(SCEN_SHEET, country=country, vehicle=vehicle, fortificant=fortificant,
                                       data_need=NEED_INT_COV):
                    target = new_value(ex, SCEN_SHEET, int_row) * fort
                    if target < baseline and not np.isclose(target, baseline):
                        problems.append((country, vehicle, fortificant, q, target, baseline))
    if problems:
        raise AssertionError(f"Intervention target below baseline coverage: {problems}")


def new_value(ex, sheet, row):
    for change in reversed(ex.changes):
        if change["sheet"] == sheet and change["row"] == row and change["status"] in ("changed", "unchanged"):
            return change["new_value"]
    return ex.get(sheet, row, "value")


def flag_remaining_placeholders(ex):
    """List placeholder rows the GF file cannot fill."""
    for sheet in [CV_SHEET, CVF_SHEET, SCEN_SHEET, "Vehicle Extraction"]:
        header = ex.header(sheet)
        if "data source" not in header:
            continue
        touched = {(c["sheet"], c["row"]) for c in ex.changes if c["status"] == "changed"}
        for row in range(2, ex.wb_values[sheet].max_row + 1):
            if norm(ex.get(sheet, row, "data source")).startswith("dummy") and (sheet, row) not in touched:
                ex.note(sheet, [row], status="still placeholder: not in GF file", gf_ref="", used=True,
                        method="no corresponding data in the GF workbook")


# --------------------------------------------------------------------------------------
# Recalculate and verify
# --------------------------------------------------------------------------------------


def recalculate(path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        return False
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            [soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(path)],
            check=True, capture_output=True,
        )
        shutil.move(str(pathlib.Path(tmp) / path.name), path)
    return True


def verify(original, updated, expected):
    """Every cell matches the original except the ones we meant to change.

    Returns (problems, recomputed). Untouched formula cells that refer to changed cells
    (e.g. the Nigeria rice Total amount, =61.2*J63) legitimately get new results; they
    are returned in ``recomputed`` for review rather than treated as errors.
    """
    old = openpyxl.load_workbook(original, data_only=True)
    old_formulas = openpyxl.load_workbook(original)
    new = openpyxl.load_workbook(updated, data_only=True)
    problems, recomputed = [], []
    for ws_old in old:
        ws_new = new[ws_old.title]
        ws_formula = old_formulas[ws_old.title]
        for row in range(1, max(ws_old.max_row, ws_new.max_row) + 1):
            for col in range(1, max(ws_old.max_column, ws_new.max_column) + 1):
                coord = ws_old.cell(row, col).coordinate
                want = expected.get((ws_old.title, coord), ws_old.cell(row, col).value)
                got = ws_new.cell(row, col).value
                if isinstance(want, (int, float)) and isinstance(got, (int, float)):
                    ok = np.isclose(want, got, rtol=1e-9, atol=1e-12)
                else:
                    ok = (want in (None, "")) and (got in (None, "")) or want == got
                if not ok:
                    original_cell = ws_formula.cell(row, col).value
                    is_formula = isinstance(original_cell, str) and original_cell.startswith("=")
                    if is_formula and (ws_old.title, coord) not in expected and got is not None:
                        recomputed.append((ws_old.title, coord, original_cell, want, got))
                    else:
                        problems.append((ws_old.title, coord, want, got))
    return problems, recomputed


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--gf", type=pathlib.Path, default=GF_FILE)
    parser.add_argument("--extraction", type=pathlib.Path, default=EXTRACTION_FILE)
    parser.add_argument("--output", type=pathlib.Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--changelog", type=pathlib.Path, default=DEFAULT_CHANGELOG)
    args = parser.parse_args()
    if args.output.resolve() == args.extraction.resolve():
        sys.exit("Refusing to overwrite the original extraction workbook; choose another --output")

    gf = GFData(args.gf)
    ex = Extraction(args.extraction)

    means = {}
    for (country, vehicle), spec in ARMS.items():
        if spec["gcap_basis"] is not None:
            any_rows = update_consumption_any(gf, ex, country, vehicle)
            means[(country, vehicle)] = update_consumption_amount(
                gf, ex, country, vehicle, spec["gcap_basis"], any_rows
            )
        update_fortifiability(gf, ex, country, vehicle)
        for fortificant in spec["fortificants"]:
            update_baseline(gf, ex, country, vehicle, fortificant)
            update_intervention(gf, ex, country, vehicle, fortificant)
    update_wheat_sds_and_u5(ex, means[("Nigeria", "Wheat")], means[("Nigeria", "Rice")])
    check_intervention_exceeds_baseline(ex)
    check_consumer_variance(ex)
    flag_remaining_placeholders(ex)

    ex.wb.calculation.fullCalcOnLoad = True
    ex.wb.save(args.output)
    log = pd.DataFrame(ex.changes)
    log.to_csv(args.changelog, index=False)

    if recalculate(args.output):
        problems, recomputed = verify(args.extraction, args.output, ex.expected)
        if problems:
            for p in problems[:30]:
                print("MISMATCH", p)
            sys.exit(f"{len(problems)} cells differ from what was intended; do not use {args.output.name}")
        print("Recalculated with LibreOffice; all untouched values match the original.")
        if recomputed:
            print("Untouched formula cells whose result changed because their inputs did:")
            for sheet, coord, formula, before, after in recomputed:
                print(f"  {sheet}!{coord} {formula}: {before} -> {after}")
    else:
        print("WARNING: LibreOffice not found. Open the new workbook in Excel and save it, "
              "or pandas will read formula cells as empty.")

    print(f"\nWrote {args.output.name} and {args.changelog.name}")
    print(log.groupby(["status"]).size().to_string())


if __name__ == "__main__":
    main()
